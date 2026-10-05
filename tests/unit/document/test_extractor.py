"""Reading pages and declared headings (spec 025, T028)."""

from types import SimpleNamespace

import pytest

from indico_assistant.services.document import extractor
from indico_assistant.services.document.extractor import UnsupportedFileTypeError, extract


def make_pdf(path, pages, outline=()):
    """A PDF with one text per page; outline entries are (title, page, level)."""
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for number, page in enumerate(pages, 1):
        for i, line in enumerate(page.split("\n")):
            c.drawString(72, 750 - 14 * i, line)
        for k, (_, on, _) in enumerate(outline):
            if on == number:
                c.bookmarkPage(f"o{k}")  # one key per entry: reportlab keys outline entries by destination
        c.showPage()
    for k, (title, _, level) in enumerate(outline):
        c.addOutlineEntry(title, f"o{k}", level=level - 1)
    c.save()
    return path


def test_pdf_pages_and_outline(tmp_path):
    path = make_pdf(
        tmp_path / "paper.pdf",
        ["Title page", "1 Introduction\nWhy it matters", "2.1 Setup\nThe rig"],
        outline=[("1 Introduction", 2, 1), ("2 Methods", 3, 1), ("2.1 Setup", 3, 2)],
    )
    doc = extract(path)
    assert len(doc.pages) == 3
    assert "Why it matters" in doc.pages[1] and "The rig" in doc.pages[2]
    assert [(h.number, h.title, h.level, h.page) for h in doc.headings] == [
        ("1", "Introduction", 1, 2),
        ("2", "Methods", 1, 3),
        ("2.1", "Setup", 2, 3),
    ]


def test_pdf_ligatures_read_as_letters(monkeypatch, tmp_path):
    """fi/ff ligature code points (what many TeX PDFs give) come out as plain letters: "different", not "dierent"."""
    pages = [SimpleNamespace(extract_text=lambda: "A diﬀerent and eﬃcient ﬁt")]
    monkeypatch.setattr("pypdf.PdfReader", lambda _: SimpleNamespace(pages=pages, outline=[]))
    doc = extract(tmp_path / "x.pdf")
    assert doc.pages == ["A different and efficient fit"]


def test_word_headings_and_page_breaks(tmp_path):
    from docx import Document

    word = Document()
    word.add_heading("Minutes", 0)
    word.add_heading("1 Budget", 1)
    word.add_paragraph("The travel budget was approved.")
    word.add_page_break()
    word.add_heading("1.1 Details", 2)
    word.add_paragraph("Flights only.")
    word.save(tmp_path / "minutes.docx")
    doc = extract(tmp_path / "minutes.docx")
    assert len(doc.pages) == 2
    assert "approved" in doc.pages[0] and "Flights only." in doc.pages[1]
    assert [(h.number, h.title, h.level, h.page) for h in doc.headings] == [
        (None, "Minutes", 1, 1),
        ("1", "Budget", 1, 1),
        ("1.1", "Details", 2, 2),
    ]


def test_powerpoint_one_page_per_slide(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for title, body, notes in [("Results", "Efficiency 93%", "Say it slowly"), ("Outlook", "Run 4", "")]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    deck.save(tmp_path / "talk.pptx")
    doc = extract(tmp_path / "talk.pptx")
    assert len(doc.pages) == 2
    assert "Efficiency 93%" in doc.pages[0] and "Notes: Say it slowly" in doc.pages[0]
    assert [(h.title, h.page) for h in doc.headings] == [("Results", 1), ("Outlook", 2)]


def test_text_and_markdown_are_one_page(tmp_path):
    (tmp_path / "a.txt").write_text("plain\ntext")
    (tmp_path / "b.md").write_text("# Plan\nintro\n## 2 Steps\ndo it")
    assert extract(tmp_path / "a.txt").pages == ["plain\ntext"]
    md = extract(tmp_path / "b.md")
    assert len(md.pages) == 1
    assert [(h.number, h.title, h.level) for h in md.headings] == [(None, "Plan", 1), ("2", "Steps", 2)]
    assert md.pages[0][md.headings[1].offset :].startswith("## 2 Steps")


def test_type_comes_from_the_filename(tmp_path):
    (tmp_path / "upload.tmp").write_text("from remote storage")
    assert extract(tmp_path / "upload.tmp", "notes.txt").pages == ["from remote storage"]
    with pytest.raises(UnsupportedFileTypeError):
        extract(tmp_path / "upload.tmp", "sheet.xlsx")


def test_control_characters_are_dropped():
    assert extractor.clean("a\x00b\tc\nd\x07") == "ab\tc\nd"
